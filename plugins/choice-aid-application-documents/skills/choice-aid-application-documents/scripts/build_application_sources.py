#!/usr/bin/env python3
"""Fill the bundled templates for native Google Sheets and Docs import.

Usage: python build_application_sources.py candidate.json output_dir
Output: JSON with resume_xlsx and career_docx absolute paths.
Requires openpyxl and python-docx. Candidate data stays outside the skill folder.
"""
import json
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Alignment
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
FONT_NAME = 'Noto Sans JP'

def use_font(cell, size=None):
    font = copy(cell.font)
    font.name = FONT_NAME
    if size is not None: font.sz = size
    cell.font = font

def resume(data, dest):
    wb = load_workbook(ROOT/'assets/resume-template.xlsm', keep_vba=False)
    ws = wb.active
    for row in ws:
        for cell in row:
            if cell.__class__.__name__ != 'MergedCell': use_font(cell)
    for key, value in {
        'C1': data.get('as_of',''), 'B2': data.get('furigana',''),
        'B3': data.get('name',''), 'B5': data.get('birth_and_age',''),
        'B6': data.get('address_furigana',''), 'B7': data.get('address',''),
        'G6': data.get('phone',''), 'F8': data.get('email',''),
        'K19': data.get('nearest_station',''), 'K18': data.get('hobbies',''),
    }.items():
        if value:
            ws[key] = value
            use_font(ws[key], 9 if key not in ('B3','B7') else 11)
    if data.get('name'): use_font(ws['B3'], 16)
    if data.get('address_furigana'):
        use_font(ws['B6'], 6)
        ws['B6'].alignment=Alignment(horizontal='left',vertical='center',shrink_to_fit=True)
    entries = data.get('resume_rows', [])
    if len(entries) > 23: raise ValueError('The source form supports 23 career rows. Review page allocation.')
    for i, (year, month, label) in enumerate(entries[:17], 14):
        if year: ws[f'A{i}'] = int(year)
        if month: ws[f'B{i}'] = int(month)
        ws[f'C{i}'] = label
        for col in 'ABC': use_font(ws[f'{col}{i}'], 8.5 if col=='C' else 9)
        ws[f'C{i}'].alignment=Alignment(horizontal='left',vertical='center',shrink_to_fit=True)
    for i, (year, month, label) in enumerate(entries[17:], 3):
        if year: ws[f'I{i}'] = int(year)
        if month: ws[f'J{i}'] = int(month)
        ws[f'K{i}'] = label
        for col in 'IJK': use_font(ws[f'{col}{i}'],8.5)
        ws[f'K{i}'].alignment=Alignment(horizontal='left',vertical='center',shrink_to_fit=True)
    for i, (year, month, label) in enumerate(data.get('qualifications', []), 10):
        if i>15: raise ValueError('Only six qualification rows available.')
        if year: ws[f'I{i}'] = int(year)
        if month: ws[f'J{i}'] = int(month)
        ws[f'K{i}'] = label
        for col in 'IJK': use_font(ws[f'{col}{i}'],8.5)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    photo_path=data.get('photo_path')
    if photo_path:
        photo=SheetImage(str(Path(photo_path).expanduser().resolve()))
        photo.width,photo.height=100,133
        # The supplied form has no photo box. Place the picture in the free
        # space immediately to its right, without covering any form cells.
        ws.add_image(photo,'N2')
        ws.print_area='A1:O30'
    else:
        ws.print_area='A1:M30'
    path = dest / f"{data['slug']}_履歴書_確認用ドラフト.xlsx"
    wb.save(path)
    return path

def set_para(p, value, size=9.5):
    p.text = value
    for run in p.runs:
        run.font.name = FONT_NAME
        run.font.size = Pt(size)

def remove(p):
    element=p._element
    element.getparent().remove(element)

def career(data, dest):
    doc=Document(ROOT/'assets/career-template.docx')
    # Google Docs conversion lays the template out more loosely than Word.
    # Preserve compact layouts without routine post-import margin edits.
    for section in doc.sections:
        section.top_margin=Pt(24)
        section.bottom_margin=Pt(24)
    ps=doc.paragraphs
    set_para(ps[2],data.get('as_of',''),9)
    set_para(ps[3],'氏名　'+data.get('name',''),10)
    set_para(ps[5],data.get('summary',''),9.5)
    target=ps[7]
    target.insert_paragraph_before('■職務経歴')
    for job in data.get('jobs',[]):
        target.insert_paragraph_before(f"{job['company']}　{job['period']}")
        target.insert_paragraph_before('職種・担当：'+job.get('role','未確認'))
        target.insert_paragraph_before(job.get('duties',''))
    set_para(ps[8],data.get('skills',''),9.5)
    set_para(ps[19],data.get('self_pr',''),9.5)
    remove(ps[9])
    for p in list(doc.paragraphs):
        if not p.text.strip(): remove(p)
    for p in doc.paragraphs:
        if p.text.startswith('■'):
            for run in p.runs: run.font.name=FONT_NAME
            p.paragraph_format.space_before=Pt(11)
            p.paragraph_format.space_after=Pt(5)
        elif p is not doc.paragraphs[0] and p.text != '以上':
            for run in p.runs:
                run.font.name=FONT_NAME
                if not run.font.size: run.font.size=Pt(9.5)
                if p.text not in (data.get('as_of',''),'氏名　'+data.get('name','')): run.bold=False
            p.paragraph_format.space_after=Pt(5)
            p.paragraph_format.line_spacing=1.15
    path=dest/f"{data['slug']}_職務経歴書_確認用ドラフト.docx"
    doc.save(path)
    return path

def main():
    if len(sys.argv)!=3: raise SystemExit(__doc__)
    data=json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
    dest=Path(sys.argv[2]).resolve()
    dest.mkdir(parents=True,exist_ok=True)
    print(json.dumps({'resume_xlsx':str(resume(data,dest)),
                      'career_docx':str(career(data,dest))},ensure_ascii=False))

if __name__=='__main__': main()
